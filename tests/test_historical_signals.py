from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document

from betexplorer_scraper.database import Database
from betexplorer_scraper.historical import HistoricalDocxImporter, compute_outcome_stats
from betexplorer_scraper.models import BookmakerOdds, DiscoveredMatch, OddsSnapshot, SnapshotQuality, TimingStatus


def _write_docx(path: Path, rows: list[list[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    table = document.add_table(rows=len(rows), cols=5)
    for row_index, row in enumerate(rows):
        for cell_index, value in enumerate(row):
            table.cell(row_index, cell_index).text = value
    document.save(path)


def _seed_match_with_required_1x2(db: Database) -> str:
    match_id = db.upsert_match(
        DiscoveredMatch(
            event_id="signal123",
            source_url="https://www.betexplorer.com/football/test/signal123/",
            league="Signal League",
            home_team="Home",
            away_team="Away",
            kickoff_time=datetime(2026, 5, 14, 20, 0),
            timing_status=TimingStatus.UPCOMING_SOON,
        )
    )
    db.save_snapshot(
        match_id,
        OddsSnapshot(
            event_id="signal123",
            market="1x2",
            captured_at=datetime(2026, 5, 14, 19, 55),
            quality_status=SnapshotQuality.COMPLETE,
            required_bookmakers=["Bwin", "Unibet"],
            bookmaker_odds=[
                BookmakerOdds("Bwin", "bwin", 2.00, 3.40, 3.00),
                BookmakerOdds("Unibet", "unibet", 2.05, 3.45, 3.10),
            ],
        ),
    )
    return match_id


def test_historical_importer_parses_docx_rows_and_score_only_rows(tmp_path: Path) -> None:
    root = tmp_path / "SAMPLE_DATABASE"
    docx_path = root / "2.00 Sample_Database ODDS" / "3.00" / "3.10.docx"
    _write_docx(
        docx_path,
        [
            ["2.00", "3.40", "3.00", "", ""],
            ["1.95", "3.35", "3.10", "2-3.", "1-1."],
            ["", "", "", "2-0.", "0-0."],
            ["not", "odds", "row", "", ""],
        ],
    )
    db = Database(tmp_path / "signals.duckdb")

    summary = HistoricalDocxImporter(db).import_roots([root])
    rows = db.list_historical_records()

    assert summary["files_imported"] == 1
    assert summary["records_imported"] == 2
    assert summary["warnings"] == 1
    assert rows[0]["dataset"] == "Odds"
    assert rows[0]["query_home_odds"] == 2.0
    assert rows[0]["query_draw_odds"] == 3.4
    assert rows[0]["query_away_odds"] == 3.0
    assert rows[0]["historical_home_odds"] == 1.95
    assert rows[0]["full_time_score"] == "2-3"
    assert rows[1]["full_time_score"] == "2-0"
    assert rows[1]["parse_status"] == "inherited_odds"


def test_historical_importer_separates_usable_odds_dataset(tmp_path: Path) -> None:
    root = tmp_path / "SAMPLE_DATABASE"
    _write_docx(
        root / "2.00 Sample_Database Usable Odds" / "3.00" / "ODDS.docx",
        [
            ["2.00", "3.40", "3.00", "", ""],
            ["2.00", "3.40", "3.00", "1-1.", "0-0."],
        ],
    )
    db = Database(tmp_path / "signals.duckdb")

    HistoricalDocxImporter(db).import_roots([root])

    assert db.list_historical_records()[0]["dataset"] == "Usable Odds"


def test_signal_stats_calculate_outcomes_totals_btts_and_double_chance() -> None:
    stats = compute_outcome_stats(["2-1", "0-0", "1-2", "2-2"])

    assert stats["sample_size"] == 4
    assert stats["home_win_pct"] == 25.0
    assert stats["draw_pct"] == 50.0
    assert stats["away_win_pct"] == 25.0
    assert stats["over_0_5_pct"] == 75.0
    assert stats["over_1_5_pct"] == 75.0
    assert stats["over_2_5_pct"] == 75.0
    assert stats["btts_pct"] == 75.0
    assert stats["double_chance_1x_pct"] == 75.0
    assert stats["double_chance_x2_pct"] == 75.0
    assert stats["double_chance_12_pct"] == 50.0


def test_database_recomputes_exact_neighbor_and_one_draw_signals(tmp_path: Path) -> None:
    db = Database(tmp_path / "signals.duckdb")
    match_id = _seed_match_with_required_1x2(db)
    db.replace_historical_records(
        "sample.docx",
        [
            {
                "dataset": "Odds",
                "source_file": "sample.docx",
                "source_home_bucket": 3.00,
                "source_away_file": 3.10,
                "query_home_odds": 2.00,
                "query_draw_odds": 3.40,
                "query_away_odds": 3.00,
                "historical_home_odds": 1.95,
                "historical_draw_odds": 3.35,
                "historical_away_odds": 3.10,
                "full_time_score": "2-1",
                "half_time_score": "1-0",
                "parse_status": "parsed",
                "parse_warning": None,
            },
            {
                "dataset": "Odds",
                "source_file": "sample.docx",
                "source_home_bucket": 3.00,
                "source_away_file": 3.10,
                "query_home_odds": 2.00,
                "query_draw_odds": 3.40,
                "query_away_odds": 3.05,
                "historical_home_odds": 2.00,
                "historical_draw_odds": 3.40,
                "historical_away_odds": 3.10,
                "full_time_score": "1-1",
                "half_time_score": "0-0",
                "parse_status": "parsed",
                "parse_warning": None,
            },
            {
                "dataset": "Usable Odds",
                "source_file": "usable.docx",
                "source_home_bucket": 3.10,
                "source_away_file": 3.10,
                "query_home_odds": 2.05,
                "query_draw_odds": 3.45,
                "query_away_odds": 3.10,
                "historical_home_odds": 2.05,
                "historical_draw_odds": 3.45,
                "historical_away_odds": 3.10,
                "full_time_score": "0-2",
                "half_time_score": "0-1",
                "parse_status": "parsed",
                "parse_warning": None,
            },
        ],
    )

    summary = db.recompute_historical_signals()
    signals = db.list_signals()
    match_signals = db.list_signals(match_id=match_id)

    assert summary["signals"] >= 4
    assert {signal["bookmaker"] for signal in signals} == {"Bwin", "Unibet"}
    assert {"exact_odds", "neighbor_odds", "one_draw"}.issubset({signal["signal_type"] for signal in signals})
    exact = next(signal for signal in signals if signal["bookmaker"] == "Bwin" and signal["signal_type"] == "exact_odds")
    assert exact["sample_size"] == 1
    assert exact["home_win_pct"] == 100.0
    assert exact["historical_scores"] == ["2-1"]
    one_draw = next(signal for signal in signals if signal["signal_type"] == "one_draw")
    assert one_draw["sample_size"] == 2
    assert one_draw["draw_pct"] == 50.0
    assert match_signals


def test_database_archives_played_matches_with_required_bookmaker_odds(tmp_path: Path) -> None:
    db = Database(tmp_path / "archive.duckdb")
    match_id = _seed_match_with_required_1x2(db)
    db.mark_result_captured(match_id, "2:1", datetime(2026, 5, 14, 22, 0))

    summary = db.archive_played_matches()
    rows = db.list_played_match_archive()

    assert summary["archived"] == 1
    assert rows[0]["event_id"] == "signal123"
    assert rows[0]["bwin_home_odds"] == 2.0
    assert rows[0]["unibet_away_odds"] == 3.1
    assert rows[0]["full_time_score"] == "2-1"
